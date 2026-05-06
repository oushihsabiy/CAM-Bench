#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""reporter.py: 结果输出 — JSON 持久化 + 终端表格 + Excel 明细。"""

from __future__ import annotations

import json
from collections import defaultdict
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path

from .difficulty_aggregator import DifficultyResult
from .evaluator import ProblemMetrics
from .solver import SolveResult


# ---------------------------------------------------------------------------
# JSON 序列化
# ---------------------------------------------------------------------------


def _serialize(obj):
    """dataclass / Path 友好的 JSON 序列化。"""
    if hasattr(obj, "__dataclass_fields__"):
        return asdict(obj)
    if isinstance(obj, Path):
        return str(obj)
    if isinstance(obj, set):
        return sorted(obj)
    raise TypeError(f"Object of type {type(obj)} is not JSON serializable")


# ---------------------------------------------------------------------------
# 写单次 SolveResult 到 LOG 目录
# ---------------------------------------------------------------------------


def save_solve_result(
    result: SolveResult,
    log_dir: Path,
    *,
    tag: str = "",
) -> Path:
    """
    保存一次求解结果到 LOG/<tag_problem_model>/。

    与现有 LOG 约定兼容：写 config.json + summary.json。
    """
    problem_stem = Path(result.problem_path).stem
    model_short = result.model.split("/")[-1]
    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    dir_name = (
        f"{ts}_{tag}_{problem_stem}_{model_short}" if tag else f"{ts}_{problem_stem}_{model_short}"
    )
    out_dir = log_dir / dir_name
    out_dir.mkdir(parents=True, exist_ok=True)

    # config.json（与现有格式兼容）
    config = {
        "plan": f"eval_{tag}" if tag else "eval",
        "problem": result.problem_path,
        "model": result.model,
        "total_blocks": result.total_blocks,
    }
    (out_dir / "config.json").write_text(
        json.dumps(config, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # summary.json
    summary = {
        "passed_count": result.passed_blocks,
        "total_blocks": result.total_blocks,
        "n_sorry_found": result.n_sorry_found,
        "all_passed": result.all_passed,
        "termination_reason": result.termination_reason,
        "mcp_enabled": result.mcp_enabled,
        "mcp_calls": result.mcp_calls,
        "mcp_failures": result.mcp_failures,
        "mcp_time_ms": result.mcp_time_ms,
        "total_attempts": result.total_attempts,
        "total_elapsed_sec": round(result.total_elapsed_sec, 2),
        "prompt_tokens": result.prompt_tokens,
        "completion_tokens": result.completion_tokens,
        "total_tokens": result.total_tokens,
        "final_verify_passed": result.final_verify_passed,
        "final_verify_issues": result.final_verify_issues,
        "block_details": [
            {
                "name": br.block_name,
                "sorry_index": br.sorry_index,
                "passed": br.passed,
                "attempts": br.attempts,
                "first_pass_turn": br.first_pass_turn,
                "elapsed_sec": round(br.total_elapsed_sec, 2),
                "prompt_tokens": br.prompt_tokens,
                "completion_tokens": br.completion_tokens,
                "total_tokens": br.total_tokens,
                "termination_reason": br.termination_reason,
            }
            for br in result.block_results
        ],
    }
    (out_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8"
    )

    # 保存每个 block 的最终代码
    for br in result.block_results:
        if br.final_code:
            fname = f"block_{br.block_index:03d}_{br.block_name}.lean"
            (out_dir / fname).write_text(br.final_code, encoding="utf-8")

    return out_dir


# ---------------------------------------------------------------------------
# 难度报告 → JSON
# ---------------------------------------------------------------------------


def _build_eval_summary(metrics: list[ProblemMetrics]) -> dict:
    """构建全局评测汇总，供 JSON 与 Markdown 报告复用。"""
    total_problems = len(metrics)
    total_runs = 0
    passed_runs = 0
    total_blocks = 0
    passed_blocks = 0
    total_attempts = 0
    total_elapsed_sec = 0.0
    total_prompt_tokens = 0
    total_completion_tokens = 0
    total_tokens = 0
    final_verify_passed_runs = 0
    final_verify_failed_runs = 0
    termination_reasons: defaultdict[str, int] = defaultdict(int)
    model_runs: defaultdict[str, int] = defaultdict(int)
    model_run_passed: defaultdict[str, int] = defaultdict(int)
    model_block_total: defaultdict[str, int] = defaultdict(int)
    model_block_passed: defaultdict[str, int] = defaultdict(int)
    difficulty_distribution: defaultdict[str, int] = defaultdict(int)

    for m in metrics:
        difficulty_distribution[m.difficulty_level] += 1
        for r in m.results:
            total_runs += 1
            total_blocks += int(r.total_blocks or 0)
            passed_blocks += int(r.passed_blocks or 0)
            total_attempts += int(r.total_attempts or 0)
            total_elapsed_sec += float(r.total_elapsed_sec or 0.0)
            total_prompt_tokens += int(r.prompt_tokens or 0)
            total_completion_tokens += int(r.completion_tokens or 0)
            total_tokens += int(r.total_tokens or 0)

            model = r.model or "unknown"
            model_runs[model] += 1
            model_block_total[model] += int(r.total_blocks or 0)
            model_block_passed[model] += int(r.passed_blocks or 0)

            if bool(r.all_passed):
                passed_runs += 1
                model_run_passed[model] += 1

            if r.final_verify_passed is True:
                final_verify_passed_runs += 1
            elif r.final_verify_passed is False:
                final_verify_failed_runs += 1

            reason = str(r.termination_reason or "unknown")
            termination_reasons[reason] += 1

    model_summary = {}
    for model in sorted(model_runs.keys()):
        mr = model_runs[model]
        mbt = model_block_total[model]
        model_summary[model] = {
            "runs": mr,
            "passed_runs": model_run_passed[model],
            "run_pass_rate": round((model_run_passed[model] / mr), 4) if mr else 0.0,
            "block_pass_rate": round((model_block_passed[model] / mbt), 4) if mbt else 0.0,
        }

    return {
        "total_problems": total_problems,
        "total_runs": total_runs,
        "passed_runs": passed_runs,
        "failed_runs": total_runs - passed_runs,
        "run_pass_rate": round((passed_runs / total_runs), 4) if total_runs else 0.0,
        "block_pass_rate": round((passed_blocks / total_blocks), 4) if total_blocks else 0.0,
        "total_blocks": total_blocks,
        "passed_blocks": passed_blocks,
        "total_attempts": total_attempts,
        "avg_attempts_per_run": round((total_attempts / total_runs), 4) if total_runs else 0.0,
        "total_elapsed_sec": round(total_elapsed_sec, 2),
        "avg_elapsed_sec_per_run": round((total_elapsed_sec / total_runs), 2) if total_runs else 0.0,
        "token_usage": {
            "prompt_tokens": total_prompt_tokens,
            "completion_tokens": total_completion_tokens,
            "total_tokens": total_tokens,
            "avg_total_tokens_per_run": round((total_tokens / total_runs), 2) if total_runs else 0.0,
        },
        "difficulty_distribution": dict(sorted(difficulty_distribution.items())),
        "termination_reasons": dict(sorted(termination_reasons.items())),
        "final_verify": {
            "passed_runs": final_verify_passed_runs,
            "failed_runs": final_verify_failed_runs,
            "unknown_runs": total_runs - final_verify_passed_runs - final_verify_failed_runs,
        },
        "per_model": model_summary,
    }


def save_eval_report(
    metrics: list[ProblemMetrics],
    output: Path,
    *,
    prior_posterior_by_problem: dict[str, dict] | None = None,
) -> None:
    """写入汇总难度报告 JSON，含每道题各项数据及每次运行/每个 block 的明细。"""
    pp_map = prior_posterior_by_problem or {}
    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "total_problems": len(metrics),
        "evaluation_summary": _build_eval_summary(metrics),
        "summary": {
            "easy": sum(1 for m in metrics if m.difficulty_level == "Easy"),
            "medium": sum(1 for m in metrics if m.difficulty_level == "Medium"),
            "hard": sum(1 for m in metrics if m.difficulty_level == "Hard"),
            "very_hard": sum(1 for m in metrics if m.difficulty_level == "Very Hard"),
        },
        "problems": [],
    }
    for m in sorted(metrics, key=lambda x: -x.difficulty_score):
        pp = pp_map.get(m.problem_path, {})
        prior = pp.get("prior", {"raw_score": 0.0, "dimensions": {}})
        posterior = pp.get("posterior", {"raw_score": 0.0, "dimensions": {}})

        # 每次运行的明细
        runs = []
        for r in m.results:
            blocks = []
            for br in r.block_results:
                blocks.append({
                    "block_name": br.block_name,
                    "block_index": br.block_index,
                    "sorry_count": br.sorry_count,
                    "passed": br.passed,
                    "attempts": br.attempts,
                    "first_pass_turn": br.first_pass_turn,
                    "elapsed_sec": round(br.total_elapsed_sec, 2),
                    "timeline": [
                        {
                            "turn": a.turn,
                            "passed": a.passed,
                            "errors": a.errors,
                            "error_types": a.error_types,
                            "mcp_context_used": a.mcp_context_used,
                            "elapsed_sec": round(a.elapsed_sec, 2),
                        }
                        for a in br.timeline
                    ],
                })
            runs.append({
                "model": r.model,
                "all_passed": r.all_passed,
                "mcp_enabled": r.mcp_enabled,
                "mcp_calls": r.mcp_calls,
                "mcp_failures": r.mcp_failures,
                "mcp_time_ms": r.mcp_time_ms,
                "passed_blocks": r.passed_blocks,
                "total_blocks": r.total_blocks,
                "total_attempts": r.total_attempts,
                "total_elapsed_sec": round(r.total_elapsed_sec, 2),
                "prompt_tokens": r.prompt_tokens,
                "completion_tokens": r.completion_tokens,
                "total_tokens": r.total_tokens,
                "termination_reason": r.termination_reason,
                "final_verify_passed": r.final_verify_passed,
                "final_verify_issues": r.final_verify_issues,
                "blocks": blocks,
            })

        report["problems"].append(
            {
                "problem": m.problem_path,
                "total_blocks": m.total_blocks,
                "difficulty_score": round(m.difficulty_score, 4),
                "difficulty_level": m.difficulty_level,
                "pass_rate": round(m.pass_rate, 4),
                "block_pass_rate": round(m.block_pass_rate, 4),
                "avg_turns_to_pass": round(m.avg_turns_to_pass, 2),
                "max_turns_to_pass": m.max_turns_to_pass,
                "avg_time_sec": round(m.avg_time_sec, 2),
                "max_time_sec": round(m.max_time_sec, 2),
                "total_prompt_tokens": m.total_prompt_tokens,
                "total_completion_tokens": m.total_completion_tokens,
                "total_tokens": m.total_tokens,
                "avg_tokens": round(m.avg_tokens, 2),
                "max_tokens": m.max_tokens,
                "error_diversity": m.error_diversity,
                "avg_proof_length": round(m.avg_proof_length, 1),
                "model_pass_rates": {k: round(v, 4) for k, v in m.model_pass_rates.items()},
                "model_gap": round(m.model_gap, 4),
                "prior": {
                    "raw_score": round(float(prior.get("raw_score", 0.0)), 4),
                    "prompt_tokens": int(prior.get("prompt_tokens", 0) or 0),
                    "completion_tokens": int(prior.get("completion_tokens", 0) or 0),
                    "total_tokens": int(prior.get("total_tokens", 0) or 0),
                    "dimensions": prior.get("dimensions", {}),
                },
                "posterior": {
                    "raw_score": round(float(posterior.get("raw_score", 0.0)), 4),
                    "prompt_tokens": int(posterior.get("prompt_tokens", 0) or 0),
                    "completion_tokens": int(posterior.get("completion_tokens", 0) or 0),
                    "total_tokens": int(posterior.get("total_tokens", 0) or 0),
                    "dimensions": posterior.get("dimensions", {}),
                },
                "total_tokens_all": (
                    int(m.total_tokens)
                    + int(prior.get("total_tokens", 0) or 0)
                    + int(posterior.get("total_tokens", 0) or 0)
                ),
                "runs": runs,
            }
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def save_eval_markdown_report(
    metrics: list[ProblemMetrics],
    output: Path,
    *,
    prior_posterior_by_problem: dict[str, dict] | None = None,
) -> None:
    """写入 Markdown 格式评测报告，便于快速浏览关键指标。"""
    summary = _build_eval_summary(metrics)
    pp_map = prior_posterior_by_problem or {}
    sorted_m = sorted(metrics, key=lambda x: -x.difficulty_score)

    lines: list[str] = []
    lines.append("# Lean Benchmark Evaluation Report")
    lines.append("")
    lines.append(f"- generated_at: {datetime.now(timezone.utc).isoformat()}")
    lines.append(f"- total_problems: {summary['total_problems']}")
    lines.append(f"- total_runs: {summary['total_runs']}")
    lines.append(f"- run_pass_rate: {summary['run_pass_rate']:.2%}")
    lines.append(f"- block_pass_rate: {summary['block_pass_rate']:.2%}")
    lines.append(f"- total_tokens: {summary['token_usage']['total_tokens']}")
    lines.append(f"- avg_tokens_per_run: {summary['token_usage']['avg_total_tokens_per_run']:.2f}")
    lines.append("")

    lines.append("## Difficulty Distribution")
    for level, count in summary["difficulty_distribution"].items():
        lines.append(f"- {level}: {count}")
    lines.append("")

    lines.append("## Termination Reasons")
    for reason, count in summary["termination_reasons"].items():
        lines.append(f"- {reason}: {count}")
    lines.append("")

    lines.append("## Per-Model Summary")
    lines.append("| Model | Runs | Passed Runs | Run Pass Rate | Block Pass Rate |")
    lines.append("| --- | ---: | ---: | ---: | ---: |")
    for model, row in summary["per_model"].items():
        lines.append(
            f"| {model} | {row['runs']} | {row['passed_runs']} | {row['run_pass_rate']:.2%} | {row['block_pass_rate']:.2%} |"
        )
    lines.append("")

    lines.append("## Problem Summary")
    lines.append("| Problem | Difficulty | Score | Run Pass Rate | Block Pass Rate | Avg Time (s) | Total Tokens |")
    lines.append("| --- | --- | ---: | ---: | ---: | ---: | ---: |")
    for m in sorted_m:
        lines.append(
            "| {problem} | {level} | {score:.4f} | {run_pr:.2%} | {block_pr:.2%} | {avg_time:.2f} | {tokens} |".format(
                problem=Path(m.problem_path).stem,
                level=m.difficulty_level,
                score=m.difficulty_score,
                run_pr=m.pass_rate,
                block_pr=m.block_pass_rate,
                avg_time=m.avg_time_sec,
                tokens=m.total_tokens,
            )
        )
    lines.append("")

    lines.append("## Run Details")
    lines.append(
        "| Problem | Model | All Passed | Passed Blocks | Total Blocks | Attempts | Time (s) | Prompt Tokens | Completion Tokens | Total Tokens | Termination | Final Verify |"
    )
    lines.append(
        "| --- | --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | ---: | --- | --- |"
    )
    for m in sorted_m:
        for r in m.results:
            lines.append(
                "| {problem} | {model} | {all_passed} | {passed_blocks} | {total_blocks} | {attempts} | {elapsed:.2f} | {pt} | {ct} | {tt} | {reason} | {fv} |".format(
                    problem=Path(m.problem_path).stem,
                    model=r.model,
                    all_passed=r.all_passed,
                    passed_blocks=r.passed_blocks,
                    total_blocks=r.total_blocks,
                    attempts=r.total_attempts,
                    elapsed=r.total_elapsed_sec,
                    pt=r.prompt_tokens,
                    ct=r.completion_tokens,
                    tt=r.total_tokens,
                    reason=r.termination_reason,
                    fv=r.final_verify_passed,
                )
            )
    lines.append("")

    lines.append("## Prior/Posterior Token Summary")
    lines.append("| Problem | Prior Tokens | Posterior Tokens | Combined Extra Tokens |")
    lines.append("| --- | ---: | ---: | ---: |")
    for m in sorted_m:
        pp = pp_map.get(m.problem_path, {})
        prior = pp.get("prior", {})
        posterior = pp.get("posterior", {})
        prior_tokens = int(prior.get("total_tokens", 0) or 0)
        posterior_tokens = int(posterior.get("total_tokens", 0) or 0)
        lines.append(
            f"| {Path(m.problem_path).stem} | {prior_tokens} | {posterior_tokens} | {prior_tokens + posterior_tokens} |"
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


# ---------------------------------------------------------------------------
# 终端打印
# ---------------------------------------------------------------------------


def print_report(metrics: list[ProblemMetrics]) -> None:
    """在终端打印按难度降序的汇总表。"""
    sorted_m = sorted(metrics, key=lambda x: -x.difficulty_score)

    print(f"\n{'=' * 110}")
    print(f"  Lean Benchmark Difficulty Report  ({len(metrics)} problems)")
    print(f"{'=' * 110}")
    print(
        f"{'Problem':<35} {'Blocks':>6} {'PassRate':>9} {'AvgTurn':>8} "
        f"{'AvgTime':>8} {'ErrDiv':>7} {'MdlGap':>7} {'Score':>7} {'Level':<10}"
    )
    print("-" * 110)

    for m in sorted_m:
        prob = Path(m.problem_path).stem[:33]
        print(
            f"{prob:<35} {m.total_blocks:>6} {m.pass_rate:>8.1%} "
            f"{m.avg_turns_to_pass:>8.1f} {m.avg_time_sec:>7.1f}s "
            f"{m.error_diversity:>7} {m.model_gap:>7.2f} "
            f"{m.difficulty_score:>7.3f} {m.difficulty_level:<10}"
        )

    dist: dict[str, int] = {}
    for m in metrics:
        dist[m.difficulty_level] = dist.get(m.difficulty_level, 0) + 1

    print("-" * 110)
    print("Distribution:", " | ".join(f"{k}: {v}" for k, v in sorted(dist.items())))
    avg_score = sum(m.difficulty_score for m in metrics) / len(metrics) if metrics else 0
    print(f"Average difficulty score: {avg_score:.3f}")
    print(f"{'=' * 110}\n")


# ---------------------------------------------------------------------------
# 评测报告 → Excel (.xlsx)
# ---------------------------------------------------------------------------


def save_eval_excel(
    metrics: list[ProblemMetrics],
    output: Path,
    *,
    prior_posterior_by_problem: dict[str, dict] | None = None,
) -> None:
    """
    写入评测明细 Excel 表。

    Sheet 1 "problems": 每道题一行，包含所有聚合指标。
    Sheet 2 "block_details": 每道题每次运行每个 block 一行。
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "Excel 导出需要 openpyxl，请先安装：pip install openpyxl"
        ) from exc

    pp_map = prior_posterior_by_problem or {}
    wb = Workbook()

    # ---- Sheet 1: problems ----
    ws = wb.active
    ws.title = "problems"

    prob_headers = [
        "problem",
        "total_blocks",
        "difficulty_score",
        "difficulty_level",
        "pass_rate",
        "block_pass_rate",
        "avg_turns_to_pass",
        "max_turns_to_pass",
        "avg_time_sec",
        "max_time_sec",
        "total_prompt_tokens",
        "total_completion_tokens",
        "total_tokens",
        "avg_tokens",
        "max_tokens",
        "error_diversity",
        "avg_proof_length",
        "model_gap",
        "num_runs",
        "model_pass_rates",
        "prior_prompt_tokens",
        "prior_completion_tokens",
        "prior_raw",
        "prior_hypothesis_count",
        "prior_sorry_block_count",
        "prior_statement_length",
        "prior_structural_complexity",
        "prior_concept_complexity",
        "prior_formalization_gap",
        "prior_type_sophistication",
        "posterior_prompt_tokens",
        "posterior_completion_tokens",
        "posterior_raw",
        "posterior_structure",
        "posterior_semantic",
        "posterior_library",
        "posterior_type",
        "posterior_search",
        "total_tokens_all",
    ]
    ws.append(prob_headers)

    sorted_m = sorted(metrics, key=lambda x: -x.difficulty_score)
    for m in sorted_m:
        pp = pp_map.get(m.problem_path, {})
        prior = pp.get("prior", {"raw_score": 0.0, "dimensions": {}})
        posterior = pp.get("posterior", {"raw_score": 0.0, "dimensions": {}})
        prior_dims = prior.get("dimensions", {})
        posterior_dims = posterior.get("dimensions", {})

        model_rates_str = "; ".join(f"{k}={v:.4f}" for k, v in m.model_pass_rates.items())
        ws.append([
            m.problem_path,
            m.total_blocks,
            round(m.difficulty_score, 4),
            m.difficulty_level,
            round(m.pass_rate, 4),
            round(m.block_pass_rate, 4),
            round(m.avg_turns_to_pass, 2),
            m.max_turns_to_pass,
            round(m.avg_time_sec, 2),
            round(m.max_time_sec, 2),
            m.total_prompt_tokens,
            m.total_completion_tokens,
            m.total_tokens,
            round(m.avg_tokens, 2),
            m.max_tokens,
            m.error_diversity,
            round(m.avg_proof_length, 1),
            round(m.model_gap, 4),
            len(m.results),
            model_rates_str,
            int(prior.get("prompt_tokens", 0) or 0),
            int(prior.get("completion_tokens", 0) or 0),
            round(float(prior.get("raw_score", 0.0)), 6),
            prior_dims.get("hypothesis_count", 0),
            prior_dims.get("sorry_block_count", 0),
            prior_dims.get("statement_length", 0),
            prior_dims.get("structural_complexity", 0),
            prior_dims.get("concept_complexity", 0),
            prior_dims.get("formalization_gap", 0),
            prior_dims.get("type_sophistication", 0),
            int(posterior.get("prompt_tokens", 0) or 0),
            int(posterior.get("completion_tokens", 0) or 0),
            round(float(posterior.get("raw_score", 0.0)), 6),
            posterior_dims.get("structure", 0),
            posterior_dims.get("semantic", 0),
            posterior_dims.get("library", 0),
            posterior_dims.get("type", posterior_dims.get("type_score", 0)),
            posterior_dims.get("search", 0),
            (
                int(m.total_tokens)
                + int(prior.get("total_tokens", 0) or 0)
                + int(posterior.get("total_tokens", 0) or 0)
            ),
        ])

    # ---- Sheet 2: block_details ----
    ws2 = wb.create_sheet("block_details")
    block_headers = [
        "problem",
        "run_index",
        "model",
        "run_all_passed",
        "run_total_elapsed_sec",
                    "run_prompt_tokens",
                    "run_completion_tokens",
                    "run_total_tokens",
        "block_index",
        "block_name",
        "sorry_count",
        "block_passed",
        "attempts",
        "first_pass_turn",
        "block_elapsed_sec",
        "error_types_seen",
    ]
    ws2.append(block_headers)

    for m in sorted_m:
        for run_idx, r in enumerate(m.results, start=1):
            for br in r.block_results:
                err_types = set()
                for a in br.timeline:
                    err_types.update(a.error_types)
                ws2.append([
                    m.problem_path,
                    run_idx,
                    r.model,
                    r.all_passed,
                    round(r.total_elapsed_sec, 2),
                    r.prompt_tokens,
                    r.completion_tokens,
                    r.total_tokens,
                    br.block_index,
                    br.block_name,
                    br.sorry_count,
                    br.passed,
                    br.attempts,
                    br.first_pass_turn,
                    round(br.total_elapsed_sec, 2),
                    "; ".join(sorted(err_types)) if err_types else "",
                ])

    # ---- 格式美化 ----
    header_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")
    for sheet in [ws, ws2]:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill = header_fill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    # 列宽
    ws.column_dimensions["A"].width = 70
    for col in "BCDEFGHIJKLMNOPQRSTUVWXYZ":
        ws.column_dimensions[col].width = 18
    ws2.column_dimensions["A"].width = 70
    ws2.column_dimensions["G"].width = 40
    ws2.column_dimensions["M"].width = 40
    for col in "BCDEFGHIJKLMNO":
        ws2.column_dimensions[col].width = 18

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


# ---------------------------------------------------------------------------
# 三阶段难度报告 → JSON
# ---------------------------------------------------------------------------


def save_difficulty_report(
    results: list[DifficultyResult],
    output: Path,
    *,
    stage_weights: dict[str, float] | None = None,
) -> None:
    """写入三阶段难度报告 JSON。"""
    from .difficulty_aggregator import STAGE_WEIGHTS as _SW

    sw = stage_weights or _SW

    report = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "total_problems": len(results),
        "stage_weights": {k: round(v, 4) for k, v in sw.items()},
        "summary": {
            "easy": sum(1 for r in results if r.difficulty_level == "Easy"),
            "medium": sum(1 for r in results if r.difficulty_level == "Medium"),
            "hard": sum(1 for r in results if r.difficulty_level == "Hard"),
            "very_hard": sum(1 for r in results if r.difficulty_level == "Very Hard"),
        },
        "problems": [],
    }

    for r in sorted(results, key=lambda x: -x.final_score):
        report["problems"].append(
            {
                "problem": r.problem_path,
                "final_score": round(r.final_score, 4),
                "difficulty_level": r.difficulty_level,
                "prior": {
                    "raw_score": round(r.prior_raw, 4),
                    "normalized": round(r.prior_norm, 4),
                    "dimensions": r.prior_dimensions,
                },
                "actual": {
                    "raw_score": round(r.actual_raw, 4),
                    "normalized": round(r.actual_norm, 4),
                    "dimensions": r.actual_dimensions,
                },
                "posterior": {
                    "raw_score": round(r.posterior_raw, 4),
                    "normalized": round(r.posterior_norm, 4),
                    "dimensions": r.posterior_dimensions,
                },
            }
        )

    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------------------------------------------------------------------------
# 三阶段难度终端打印
# ---------------------------------------------------------------------------


def print_difficulty_report(results: list[DifficultyResult]) -> None:
    """在终端打印三阶段难度汇总表。"""
    sorted_r = sorted(results, key=lambda x: -x.final_score)

    print(f"\n{'=' * 120}")
    print(f"  Lean Benchmark Three-Stage Difficulty Report  ({len(results)} problems)")
    print(f"{'=' * 120}")
    print(
        f"{'Problem':<35} {'Prior':>8} {'PrNorm':>7} "
        f"{'Actual':>8} {'AcNorm':>7} "
        f"{'Post':>8} {'PoNorm':>7} "
        f"{'Final':>7} {'Level':<10}"
    )
    print("-" * 120)

    for r in sorted_r:
        prob = Path(r.problem_path).stem[:33]
        print(
            f"{prob:<35} {r.prior_raw:>8.3f} {r.prior_norm:>7.3f} "
            f"{r.actual_raw:>8.3f} {r.actual_norm:>7.3f} "
            f"{r.posterior_raw:>8.3f} {r.posterior_norm:>7.3f} "
            f"{r.final_score:>7.3f} {r.difficulty_level:<10}"
        )

    dist: dict[str, int] = {}
    for r in results:
        dist[r.difficulty_level] = dist.get(r.difficulty_level, 0) + 1

    print("-" * 120)
    print("Distribution:", " | ".join(f"{k}: {v}" for k, v in sorted(dist.items())))
    avg = sum(r.final_score for r in results) / len(results) if results else 0
    print(f"Average final difficulty: {avg:.3f}")
    print(f"{'=' * 120}\n")


# ---------------------------------------------------------------------------
# 三阶段难度报告 → Excel (.xlsx)
# ---------------------------------------------------------------------------


def save_difficulty_excel(
    results: list[DifficultyResult],
    output: Path,
    *,
    stage_weights: dict[str, float] | None = None,
) -> None:
    """
    写入 Excel 明细表。

    包含：
    1) 终端表格中所有字段
    2) 每题 prior / posterior 每一个维度分数
    3) actual 各维度分数（便于排查）
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError(
            "Excel 导出需要 openpyxl，请先安装：pip install openpyxl"
        ) from exc

    from .difficulty_aggregator import STAGE_WEIGHTS as _SW

    sw = stage_weights or _SW
    sorted_r = sorted(results, key=lambda x: -x.final_score)

    wb = Workbook()
    ws = wb.active
    ws.title = "difficulty_details"

    headers = [
        "problem",
        "prior_raw",
        "prior_norm",
        "actual_raw",
        "actual_norm",
        "posterior_raw",
        "posterior_norm",
        "final_score",
        "difficulty_level",
        "prior_hypothesis_count",
        "prior_sorry_block_count",
        "prior_statement_length",
        "prior_structural_complexity",
        "prior_concept_complexity",
        "prior_formalization_gap",
        "prior_type_sophistication",
        "posterior_structure",
        "posterior_semantic",
        "posterior_library",
        "posterior_type",
        "posterior_search",
        "actual_pass_rate_inv",
        "actual_turns_norm",
        "actual_time_norm",
        "actual_error_div_norm",
        "actual_model_gap",
    ]

    ws.append(headers)

    for r in sorted_r:
        prior_dims = r.prior_dimensions or {}
        post_dims = r.posterior_dimensions or {}
        actual_dims = r.actual_dimensions or {}

        ws.append(
            [
                r.problem_path,
                round(r.prior_raw, 6),
                round(r.prior_norm, 6),
                round(r.actual_raw, 6),
                round(r.actual_norm, 6),
                round(r.posterior_raw, 6),
                round(r.posterior_norm, 6),
                round(r.final_score, 6),
                r.difficulty_level,
                prior_dims.get("hypothesis_count", 0),
                prior_dims.get("sorry_block_count", 0),
                prior_dims.get("statement_length", 0),
                prior_dims.get("structural_complexity", 0),
                prior_dims.get("concept_complexity", 0),
                prior_dims.get("formalization_gap", 0),
                prior_dims.get("type_sophistication", 0),
                post_dims.get("structure", 0),
                post_dims.get("semantic", 0),
                post_dims.get("library", 0),
                post_dims.get("type", post_dims.get("type_score", 0)),
                post_dims.get("search", 0),
                actual_dims.get("pass_rate_inv", 0),
                actual_dims.get("turns_norm", 0),
                actual_dims.get("time_norm", 0),
                actual_dims.get("error_div_norm", 0),
                actual_dims.get("model_gap", 0),
            ]
        )

    header_fill = PatternFill(fill_type="solid", fgColor="D9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.fill = header_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    width_map = {
        "A": 70,
        "I": 14,
    }
    for col in "BCDEFGHJKLMNOPQRSTUVWXYZ":
        width_map[col] = 16
    for col, width in width_map.items():
        ws.column_dimensions[col].width = width

    summary = wb.create_sheet("summary")
    summary.append(["generated_at", datetime.now(timezone.utc).isoformat()])
    summary.append(["total_problems", len(results)])
    summary.append(["stage_weight_prior", sw.get("prior", 0.0)])
    summary.append(["stage_weight_actual", sw.get("actual", 0.0)])
    summary.append(["stage_weight_posterior", sw.get("posterior", 0.0)])
    summary.append(["easy", sum(1 for r in results if r.difficulty_level == "Easy")])
    summary.append(["medium", sum(1 for r in results if r.difficulty_level == "Medium")])
    summary.append(["hard", sum(1 for r in results if r.difficulty_level == "Hard")])
    summary.append(["very_hard", sum(1 for r in results if r.difficulty_level == "Very Hard")])

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
